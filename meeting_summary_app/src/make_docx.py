"""המרת סיכום (Markdown פשוט) למסמך Word בעברית: David 12, ימין לשמאל, יישור לשני הצדדים,
ריווח שורה וחצי, כותרות גדולות בכחול כהה, טבלאות מימין לשמאל.

שימוש:  python make_docx.py <summary.md> <output.docx>

תחביר נתמך:  # כותרת | ## כותרת סעיף | ### תת-כותרת | - תבליט | 1. מספור
             ☐ משימה | שורות טבלה |a|b| | **מודגש**
"""
import re
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "David"
FOOTER_TEXT = "מסכם הישיבות של ליאור אגאי"
BLACK = RGBColor(0, 0, 0)
DARK_BLUE = RGBColor(0x1F, 0x38, 0x64)
HEADING_SIZES = {1: 18, 2: 15, 3: 13}  # טקסט רגיל: 12

# סדר הרכיבים שנדרש לפי תקן OOXML (Word עלול לדווח על תוכן לא קריא אם הסדר שגוי)
ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
           "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
           "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
           "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs "
           "em lang eastAsianLayout specVanish oMath",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange",
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType "
              "pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings "
              "sectPrChange",
}


def reorder(el):
    order = {qn("w:" + t): i for i, t in enumerate(ORDER[el.tag.split("}")[1]].split())}
    for child in sorted(el, key=lambda c: order.get(c.tag, len(order))):
        el.append(child)  # append מעביר את הרכיב לסוף, כך שהסדר נבנה מחדש


def set_rtl_par(p):
    # כל פסקה: ימין לשמאל, מיושרת לשני הצדדים, ריווח שורה וחצי
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.append(OxmlElement("w:bidi"))
    reorder(pPr)


def style_run(run, size=12, bold=False, color=BLACK):
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rPr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    rPr.append(OxmlElement("w:rtl"))
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(size * 2))
    rPr.append(szCs)
    if bold:
        rPr.append(OxmlElement("w:bCs"))
    reorder(rPr)


def add_text(p, text, size=12, bold=False, color=BLACK):
    # **מודגש** בתוך שורה
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            style_run(p.add_run(part), size, bold or i % 2 == 1, color)


def add_par(doc, text, style=None, size=12, bold=False, color=BLACK):
    p = doc.add_paragraph(style=style)
    set_rtl_par(p)
    add_text(p, text, size, bold, color)
    return p


def add_heading(doc, text, level):
    p = add_par(doc, text.replace("**", ""), size=HEADING_SIZES[level], bold=True, color=DARK_BLUE)
    p.paragraph_format.keep_with_next = True
    reorder(p._p.get_or_add_pPr())
    return p


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    # טבלה מימין לשמאל: העמודה הראשונה מימין, והטבלה צמודה לשוליים הימניים
    tblPr = table._tbl.tblPr
    tblPr.append(OxmlElement("w:bidiVisual"))
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "right")
    tblPr.append(jc)
    reorder(tblPr)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            p = table.cell(r, c).paragraphs[0]
            set_rtl_par(p)
            add_text(p, val, 12, bold=(r == 0))
    doc.add_paragraph()


def add_field(p, instr, size=10):
    # שדה Word (PAGE / NUMPAGES) שמתעדכן אוטומטית
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        run = p.add_run()
        style_run(run, size)
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            run.text = text


def set_alignment(p, align):
    # שינוי יישור אחרי set_rtl_par - שומרים על סדר הרכיבים
    p.alignment = align
    reorder(p._p.get_or_add_pPr())


def add_footer(section):
    p = section.footer.paragraphs[0]
    set_rtl_par(p)
    set_alignment(p, WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, f"{FOOTER_TEXT}   |   עמוד ", 10)
    add_field(p, "PAGE")
    add_text(p, " מתוך ", 10)
    add_field(p, "NUMPAGES")


def main(src, dst):
    doc = Document()
    doc.sections[0]._sectPr.append(OxmlElement("w:bidi"))
    reorder(doc.sections[0]._sectPr)
    add_footer(doc.sections[0])
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal._element.get_or_add_pPr().append(OxmlElement("w:bidi"))
    reorder(normal._element.get_or_add_pPr())

    table_rows = []
    lines = open(src, encoding="utf-8-sig").read().splitlines()
    for line in lines + [""]:
        s = line.strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                table_rows.append(cells)
            continue
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []
        if not s:
            continue
        if s.startswith("# "):
            add_heading(doc, s[2:], 1)
        elif s.startswith("## "):
            add_heading(doc, s[3:], 2)
        elif s.startswith("### "):
            add_heading(doc, s[4:], 3)
        elif re.fullmatch(r"\*\*[^*]+\*\*:?", s):
            # שורה שכולה מודגשת = כותרת
            add_heading(doc, s.rstrip(":").strip("*").strip() + (":" if s.endswith(":") else ""), 3)
        elif s.startswith("- "):
            add_par(doc, s[2:], style="List Bullet")
        elif re.match(r"\d+\.\s", s):
            add_par(doc, re.sub(r"^\d+\.\s+", "", s), style="List Number")
        else:
            add_par(doc, s)

    doc.save(dst)
    print(dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
