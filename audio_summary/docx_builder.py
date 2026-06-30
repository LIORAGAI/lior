"""בניית מסמך Word מעוצב RTL (גופן David, יישור ימין) מתוך טקסט Markdown מובנה."""
import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_NAME = "David"
FONT_SIZE = Pt(12)


def _set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _set_run_font(run):
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), FONT_NAME)
    rFonts.set(qn("w:ascii"), FONT_NAME)
    rFonts.set(qn("w:hAnsi"), FONT_NAME)
    rtl = OxmlElement("w:rtl")
    rPr.append(rtl)


def _add_paragraph(doc, text, bold=False, size=None, style=None):
    p = doc.add_paragraph(style=style)
    _set_rtl(p)
    run = p.add_run(text)
    _set_run_font(run)
    run.bold = bold
    if size:
        run.font.size = size
    return p


def _set_table_rtl(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tblPr.append(bidi)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT


def _add_table(doc, header, rows):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    _set_table_rtl(table)
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header):
        p = hdr_cells[i].paragraphs[0]
        _set_rtl(p)
        run = p.add_run(h)
        _set_run_font(run)
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            _set_rtl(p)
            run = p.add_run(val)
            _set_run_font(run)
    return table


def _parse_md_table(lines, start_idx):
    """מפענח טבלת Markdown החל מ-start_idx. מחזיר (header, rows, next_idx)."""
    header_line = lines[start_idx]
    header = [c.strip() for c in header_line.strip().strip("|").split("|")]
    idx = start_idx + 1
    if idx < len(lines) and re.match(r"^\s*\|?[\s:|-]+\|?\s*$", lines[idx]):
        idx += 1
    rows = []
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        row = [c.strip() for c in lines[idx].strip().strip("|").split("|")]
        rows.append(row)
        idx += 1
    return header, rows, idx


def build_docx(markdown_text: str, output_path: str):
    doc = Document()

    section = doc.sections[0]
    sectPr = section._sectPr
    bidi = OxmlElement("w:bidi")
    sectPr.append(bidi)

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = FONT_SIZE
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), FONT_NAME)

    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("# "):
            _add_paragraph(doc, stripped[2:].strip(), bold=True, size=Pt(18))
        elif stripped.startswith("## "):
            _add_paragraph(doc, stripped[3:].strip(), bold=True, size=Pt(14))
        elif stripped.startswith("- [ ]") or stripped.startswith("☐"):
            text = stripped.replace("- [ ]", "").strip()
            if not text.startswith("☐"):
                text = f"☐ {text}"
            _add_paragraph(doc, text)
        elif stripped.startswith("|"):
            header, rows, next_i = _parse_md_table(lines, i)
            _add_table(doc, header, rows)
            i = next_i
            continue
        elif stripped.startswith("**") and stripped.endswith("**") and len(stripped) > 4:
            _add_paragraph(doc, stripped.strip("*"), bold=True)
        else:
            _add_paragraph(doc, stripped)

        i += 1

    doc.save(output_path)
